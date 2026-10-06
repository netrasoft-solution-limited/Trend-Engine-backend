"""Word. PRD §14 names DOCX as one of the two required export formats.

Built imperatively rather than from the HTML, because there is no honest way to
convert one to the other without either a converter dependency or a second
layout implementation that drifts. Both renderers consume the same parsed
blocks, so the structure cannot diverge even though the styling is written
twice.

DETERMINISM TAKES REAL WORK HERE, and it is the reason this file is longer than
`markdown.py`. A .docx is a zip, and a zip records a modification time per entry
from the clock — so the same document saved twice produces different bytes for
no reason anyone would ever see. `ExportArtifact.sha256` exists so a file can be
re-rendered and proved identical to what a client received; without the
normalisation below that column would be noise that looks like evidence.

Two sources of drift, both pinned:

  · Core properties — python-docx stamps `created` and `modified` with
    `datetime.now()`.
  · Zip entry timestamps — written from the local clock at save time.
"""
from __future__ import annotations

import io
import zipfile
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .document import ExportDocument

MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXTENSION = "docx"

#: Any fixed instant. The zip format cannot represent anything before 1980, so
#: the usual Unix epoch is not available.
FIXED_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_DATETIME = datetime(1980, 1, 1)

INK = RGBColor(0x1C, 0x1B, 0x19)
MUTED = RGBColor(0x6B, 0x68, 0x64)
WARNING = RGBColor(0xA3, 0x2B, 0x20)


def render(document: ExportDocument) -> bytes:
    doc = Document()
    _set_base_style(doc)
    _pin_core_properties(doc)

    if document.is_draft:
        notice = doc.add_paragraph()
        notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = notice.add_run(document.notice)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = WARNING

    title = doc.add_paragraph()
    title_run = title.add_run(document.title)
    title_run.bold = True
    title_run.font.size = Pt(19)
    title_run.font.color.rgb = INK

    meta_parts = [document.type_label, f"Prepared for {document.organization_name}"]
    if document.dated:
        meta_parts.append(document.dated.strftime("%-d %B %Y"))
    meta = doc.add_paragraph()
    meta_run = meta.add_run(" · ".join(part for part in meta_parts if part))
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = MUTED

    if document.summary:
        summary = doc.add_paragraph()
        summary_run = summary.add_run(document.summary)
        summary_run.font.size = Pt(11.5)

    for section in document.sections:
        heading = doc.add_paragraph()
        heading_run = heading.add_run(section.heading)
        heading_run.bold = True
        heading_run.font.size = Pt(12)
        # Word's own "keep with next", so a heading never sits alone at the
        # foot of a page — the same rule the print stylesheet sets for PDF.
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(14)

        for block in section.blocks:
            if block.is_bullets:
                for item in block.items:
                    # "List Bullet" is a built-in style, so the list survives
                    # into Google Docs and Pages rather than relying on a style
                    # this file would have to define.
                    doc.add_paragraph(item, style="List Bullet")
            else:
                doc.add_paragraph(block.text)

    if not document.provenance.is_empty:
        parts = []
        if document.provenance.client_profile_version:
            parts.append(f"client profile {document.provenance.client_profile_version}")
        if document.provenance.domain_pack_version:
            parts.append(f"domain pack {document.provenance.domain_pack_version}")
        footer = doc.add_paragraph()
        footer.paragraph_format.space_before = Pt(18)
        footer_run = footer.add_run(f"Built from {', '.join(parts)}.")
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = MUTED

    buffer = io.BytesIO()
    doc.save(buffer)
    return _normalise(buffer.getvalue())


def _set_base_style(doc) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)


def _pin_core_properties(doc) -> None:
    """python-docx stamps these from the clock. Left alone, two renders of one
    document differ in `docProps/core.xml` and nowhere else."""
    properties = doc.core_properties
    properties.created = FIXED_DATETIME
    properties.modified = FIXED_DATETIME
    properties.last_modified_by = ""
    properties.revision = 1


def _normalise(raw: bytes) -> bytes:
    """Rewrite the archive with fixed entry timestamps.

    Order, compression and content are preserved exactly; only the clock is
    removed. Done as a post-pass rather than by patching zipfile, because
    monkey-patching a stdlib module to make one caller deterministic is the kind
    of fix that breaks something unrelated a year later.
    """
    source = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME)
            entry.compress_type = info.compress_type
            entry.external_attr = info.external_attr
            entry.internal_attr = info.internal_attr
            entry.create_system = info.create_system
            target.writestr(entry, source.read(info.filename))
    return out.getvalue()
