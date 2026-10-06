"""Rendering an output into something a client can be sent.

With no portal, a file IS the delivery, so these renderers are the last mile
rather than a convenience. Three things are worth defending:

  · THE SHARED PARSE. One body must produce the same structure in every format.
    If each renderer reads the `\\n\\n` and `·` conventions independently, a
    bullet list silently becomes a run-on paragraph in Word and reads as sloppy
    writing rather than as a bug.
  · DETERMINISM. `ExportArtifact.sha256` is only worth storing if re-rendering
    reproduces the bytes. Otherwise it identifies a blob nobody has while
    looking like evidence.
  · THE DRAFT WATERMARK. This machine can turn an unapproved version into a
    client-ready PDF. The gate exists because approved is not visible; the
    watermark is what is left of that backstop once a file can be forwarded.
"""
from __future__ import annotations

from datetime import date

import pytest

from apps.outputs import exports
from apps.outputs.exports import text
from apps.outputs.exports.document import Provenance, document_from_body

# The real shapes `apps/outputs/drafting.py` emits: paragraphs separated by a
# blank line, and a lead-in followed by `·` bullets in the methodology section.
BODY = [
    {
        "heading": "What the evidence says it does",
        "text": (
            "Magnesium glycinate improves sleep onset. “I take it an hour before "
            "bed” — Dr Example, “Sleep and recovery” at 12:04.\n\n"
            "A second claim, in its own paragraph."
        ),
    },
    {
        "heading": "What this is based on",
        "text": "Sources (2):\n· Sleep and recovery — Dr Example\n· The long game — Someone Else",
    },
]


def doc(**overrides):
    defaults = dict(
        title="Magnesium — what moved in October",
        type_label="Trend intelligence brief",
        organization_name="Jarrow Formulas",
        summary="Three things worth knowing.",
        body=BODY,
        is_draft=False,
        dated=date(2026, 10, 5),
        provenance=Provenance(version_number=2, client_profile_version="v3"),
    )
    return document_from_body(**{**defaults, **overrides})


# ── The shared parse ────────────────────────────────────────────────────────


def test_a_blank_line_separates_paragraphs():
    blocks = text.blocks("First paragraph.\n\nSecond paragraph.")

    assert [b.text for b in blocks] == ["First paragraph.", "Second paragraph."]


def test_bullet_lines_become_one_list():
    blocks = text.blocks("· Alpha\n· Beta\n· Gamma")

    assert len(blocks) == 1
    assert blocks[0].is_bullets
    assert blocks[0].items == ("Alpha", "Beta", "Gamma")


def test_a_lead_in_before_a_list_stays_its_own_paragraph():
    """"Sources (2):" must not be swallowed into the first bullet, which is what
    a naive line-by-line parse does."""
    blocks = text.blocks("Sources (2):\n· Alpha\n· Beta")

    assert len(blocks) == 2
    assert blocks[0].text == "Sources (2):"
    assert blocks[1].items == ("Alpha", "Beta")


def test_prose_wrapped_over_several_lines_is_one_paragraph():
    blocks = text.blocks("A sentence that was\nwrapped across lines.")

    assert len(blocks) == 1
    assert blocks[0].text == "A sentence that was wrapped across lines."


def test_a_list_followed_by_prose_splits_again():
    blocks = text.blocks("· Alpha\n· Beta\nAnd then a closing note.")

    assert [b.is_bullets for b in blocks] == [True, False]


def test_anything_unrecognised_is_still_a_paragraph():
    """Total by design: a malformed body must be an ugly document, never an
    export outage."""
    assert text.blocks("---").__len__() == 1
    assert text.blocks("") == ()
    assert text.blocks("   \n\n  ") == ()


def test_windows_line_endings_parse_the_same():
    assert text.blocks("A.\r\n\r\nB.") == text.blocks("A.\n\nB.")


def test_every_format_sees_the_same_structure():
    """The point of the shared parse. A bullet list has to be a list in all of
    them, not a paragraph in some."""
    document = doc()
    methodology = document.sections[1]

    assert methodology.blocks[1].is_bullets
    assert methodology.blocks[1].items == (
        "Sleep and recovery — Dr Example",
        "The long game — Someone Else",
    )


# ── Determinism ─────────────────────────────────────────────────────────────


@pytest.mark.parametrize("fmt", ["markdown", "html", "csv"])
def test_rendering_twice_produces_identical_bytes(fmt):
    """`ExportArtifact.sha256` is only meaningful if this holds. A "generated
    at" line anywhere would break it and make the hash worthless."""
    first = exports.render(doc(), fmt)
    second = exports.render(doc(), fmt)

    assert first.content == second.content
    assert first.sha256() == second.sha256()


@pytest.mark.parametrize("fmt", ["markdown", "html", "csv"])
def test_no_clock_reading_reaches_any_format(fmt):
    """The distinction determinism rests on: a publication date is a fact about
    the document and identical on every render; the time of rendering is not,
    and one line of it would make every sha256 differ."""
    from datetime import datetime

    body = exports.render(doc(), fmt).content.decode()

    assert datetime.now().strftime("%H:%M") not in body


@pytest.mark.parametrize("fmt", ["markdown", "html"])
def test_the_formats_with_a_header_carry_the_publication_date(fmt):
    """CSV is excluded on purpose: it is section,paragraph and has nowhere to
    put metadata, which is the honest shape rather than a missing feature."""
    body = exports.render(doc(), fmt).content.decode()

    assert "2026" in body


@pytest.mark.parametrize("fmt", ["markdown", "html", "csv"])
def test_the_filename_is_stable_and_safe(fmt):
    rendered = exports.render(doc(), fmt)

    assert rendered.filename.startswith("magnesium-what-moved-in-october")
    assert " " not in rendered.filename


def test_a_title_of_only_punctuation_still_produces_a_filename():
    rendered = exports.render(doc(title="!!!"), "markdown")

    assert rendered.filename == "output.md"


def test_an_accented_title_becomes_an_ascii_filename():
    rendered = exports.render(doc(title="Café résumé"), "markdown")

    assert rendered.filename == "cafe-resume.md"


# ── The draft watermark ─────────────────────────────────────────────────────


@pytest.mark.parametrize("fmt", ["markdown", "html", "csv"])
def test_an_unpublished_document_is_watermarked_in_every_format(fmt):
    """This machine can hand an operator a client-ready file built from
    unapproved content. The gate's backstop is this notice."""
    body = exports.render(doc(is_draft=True), fmt).content.decode()

    assert exports.DRAFT_NOTICE in body


@pytest.mark.parametrize("fmt", ["markdown", "html", "csv"])
def test_a_published_document_carries_no_watermark(fmt):
    body = exports.render(doc(is_draft=False), fmt).content.decode()

    assert exports.DRAFT_NOTICE not in body


def test_the_watermark_is_at_the_top_not_the_bottom():
    """Someone who forwards the first screenful has forwarded the warning."""
    body = exports.render(doc(is_draft=True), "markdown").content.decode()

    assert body.index(exports.DRAFT_NOTICE) < body.index("Magnesium")


# ── What each format must actually contain ──────────────────────────────────


def test_markdown_renders_headings_and_bullets():
    body = exports.render(doc(), "markdown").content.decode()

    assert "# Magnesium — what moved in October" in body
    assert "## What the evidence says it does" in body
    assert "- Sleep and recovery — Dr Example" in body


def test_markdown_ends_with_exactly_one_newline():
    content = exports.render(doc(), "markdown").content

    assert content.endswith(b"\n")
    assert not content.endswith(b"\n\n")


def test_html_escapes_client_content():
    """A claim quoting someone who said "<" must not produce markup."""
    body = exports.render(
        doc(body=[{"heading": "H", "text": "He said <script>alert(1)</script> on air."}]),
        "html",
    ).content.decode()

    assert "<script>alert(1)</script>" not in body
    assert "&lt;script&gt;" in body


def test_html_does_not_inherit_the_operator_shell():
    """A document going to a client must not carry the nav, the sign-out button
    or an operator's email address."""
    body = exports.render(doc(), "html").content.decode()

    assert "Sign out" not in body
    assert "prefers-color-scheme" not in body
    assert "color-mix" not in body


def test_html_carries_print_rules_so_the_pdf_paginates():
    body = exports.render(doc(), "html").content.decode()

    assert "@page" in body
    assert "page-break-after: avoid" in body, "a heading alone at a page foot"


def test_csv_is_one_row_per_paragraph_with_its_section():
    body = exports.render(doc(), "csv").content.decode()
    rows = [r for r in body.split("\r\n") if r]

    assert rows[0] == "section,paragraph"
    assert any(r.startswith("What this is based on,Sleep and recovery") for r in rows)


def test_csv_uses_rfc_line_endings_regardless_of_platform():
    """Otherwise the bytes follow the machine that rendered them and the sha256
    is meaningless."""
    content = exports.render(doc(), "csv").content

    assert b"\r\n" in content


def test_csv_quotes_a_paragraph_containing_a_comma():
    body = exports.render(
        doc(body=[{"heading": "H", "text": "One, two, three."}]), "csv"
    ).content.decode()

    assert '"One, two, three."' in body


# ── Provenance ──────────────────────────────────────────────────────────────


def test_the_document_names_the_profile_it_was_built_from():
    """PRD §8: every output records exactly one tenant, one client-profile
    version and one domain-pack version."""
    for fmt in ("markdown", "html"):
        body = exports.render(doc(), fmt).content.decode()
        assert "client profile v3" in body


def test_an_output_with_no_profile_says_nothing_rather_than_saying_none():
    body = exports.render(doc(provenance=Provenance()), "markdown").content.decode()

    assert "Built from" not in body


# ── The registry ────────────────────────────────────────────────────────────


def test_the_three_dependency_free_formats_are_always_available():
    available = exports.available()

    assert {"markdown", "html", "csv"} <= set(available)


def test_an_unknown_format_is_refused_by_name():
    with pytest.raises(exports.ExportError, match="not an available export format"):
        exports.render(doc(), "wordperfect")


def test_rendered_carries_what_an_artifact_row_needs():
    rendered = exports.render(doc(), "markdown")

    assert rendered.byte_size == len(rendered.content)
    assert len(rendered.sha256()) == 64
    assert rendered.mimetype.startswith("text/markdown")


# ── DOCX and PDF ────────────────────────────────────────────────────────────
#
# Both need a third-party package, and PDF additionally needs system libraries
# and a font. They skip rather than fail when absent, so a developer without
# them still gets a meaningful run — but CI builds the image, where they must
# pass. See `deploy/Dockerfile`.

requires_docx = pytest.mark.skipif(
    "docx" not in exports.available(), reason="python-docx is not installed"
)
requires_pdf = pytest.mark.skipif(
    "pdf" not in exports.available(), reason="weasyprint and its system libraries are not installed"
)


@requires_docx
def test_a_docx_renders_twice_to_identical_bytes():
    """The hardest determinism case: a .docx is a zip, and a zip stamps every
    entry from the clock. Without the normalisation in `docx.py` this fails,
    and `ExportArtifact.sha256` becomes noise that looks like evidence."""
    first = exports.render(doc(), "docx")
    second = exports.render(doc(), "docx")

    assert first.content == second.content


@requires_docx
def test_every_zip_entry_timestamp_is_pinned():
    import io
    import zipfile

    from apps.outputs.exports.docx import FIXED_TIME

    archive = zipfile.ZipFile(io.BytesIO(exports.render(doc(), "docx").content))

    assert all(entry.date_time == FIXED_TIME for entry in archive.infolist())


@requires_docx
def test_the_docx_core_properties_do_not_carry_the_clock():
    """python-docx stamps created/modified with `datetime.now()`, which would
    differ between renders in `docProps/core.xml` and nowhere else."""
    import io
    import zipfile
    from datetime import datetime

    archive = zipfile.ZipFile(io.BytesIO(exports.render(doc(), "docx").content))
    core = archive.read("docProps/core.xml").decode()

    assert str(datetime.now().year) not in core


@requires_docx
def test_the_docx_is_a_document_word_can_open():
    """Round-tripping through the library proves the package is well-formed,
    which a zip-structure check alone does not."""
    import io

    from docx import Document

    opened = Document(io.BytesIO(exports.render(doc(), "docx").content))
    text_of = [p.text for p in opened.paragraphs if p.text.strip()]

    assert "Magnesium — what moved in October" in text_of


@requires_docx
def test_bullets_become_real_list_paragraphs_not_prose():
    """The failure the shared parse exists to prevent: a list flattened into one
    run-on paragraph reads as sloppy writing rather than as a bug."""
    import io

    from docx import Document

    opened = Document(io.BytesIO(exports.render(doc(), "docx").content))
    bullets = [p.text for p in opened.paragraphs if p.style.name == "List Bullet"]

    assert bullets == ["Sleep and recovery — Dr Example", "The long game — Someone Else"]


@requires_docx
def test_the_docx_carries_the_draft_notice():
    import io

    from docx import Document

    opened = Document(io.BytesIO(exports.render(doc(is_draft=True), "docx").content))

    assert any(exports.DRAFT_NOTICE in p.text for p in opened.paragraphs)


@requires_pdf
def test_a_pdf_renders_twice_to_identical_bytes():
    """A PDF embeds /CreationDate, /ModDate and a document /ID, all of which
    default to the clock or to randomness."""
    first = exports.render(doc(), "pdf")
    second = exports.render(doc(), "pdf")

    assert first.content == second.content


@requires_pdf
def test_the_pdf_is_a_pdf_with_pages():
    rendered = exports.render(doc(), "pdf")

    assert rendered.content.startswith(b"%PDF-")
    assert b"/Type /Page" in rendered.content or b"/Type/Page" in rendered.content


@requires_pdf
def test_the_pdf_embeds_a_font_covering_the_punctuation_this_content_uses():
    """Curly quotes and `·` are everywhere in this content. In a container with
    no font package they render as empty boxes — the PDF succeeds and looks
    broken, which is the harder fault to notice. An embedded font descriptor is
    the evidence that something was available to draw them."""
    content = exports.render(doc(), "pdf").content

    assert b"/FontFile" in content or b"/BaseFont" in content


@requires_pdf
def test_a_draft_pdf_differs_from_the_published_one():
    """The watermark is drawn, so the bytes must differ. Reading text back out
    of a PDF would need another dependency to assert something this already
    shows."""
    assert exports.render(doc(is_draft=True), "pdf").content != exports.render(
        doc(is_draft=False), "pdf"
    ).content


def test_the_pdf_renderer_pins_the_font_subsetter_clock():
    """Runs even where weasyprint cannot load, because this is the setting the
    whole of PDF determinism rests on and it is easy to delete by accident.

    Two renders differed by a few bytes inside the embedded font program — the
    TrueType `head` table, which the subsetter stamps from the clock. Nothing in
    the PDF metadata explains it, so the next person to see sha256 values drift
    would reasonably spend a day on /CreationDate.
    """
    import os

    assert os.environ.get("SOURCE_DATE_EPOCH") == exports.FIXED_EPOCH
